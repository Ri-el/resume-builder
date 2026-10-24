"""Contract tests for the supplied format, content preservation, and fitting policy."""

import contextlib
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches

from scripts import build_resume as builder
from scripts.resume_content import ContentError, content_fragments, load_content, validate_content
from scripts.resume_layout import build_document, fitting_stages, load_profile

ROOT = Path(__file__).resolve().parents[1]


def sample():
    return json.loads((ROOT / "examples/candidate.json").read_text(encoding="utf-8"))


class ContentTests(unittest.TestCase):
    def test_validation_does_not_modify_candidate_content(self):
        candidate = sample()
        before = copy.deepcopy(candidate)
        self.assertIs(validate_content(candidate), candidate)
        self.assertEqual(candidate, before)

    def test_unknown_fields_and_missing_fields_are_not_silently_lost(self):
        for change in (lambda c: c.update(extra="unmapped fact"), lambda c: c.pop("name"),
                       lambda c: c["experience"][0].update(achievement="unmapped fact")):
            with self.subTest(change=change):
                candidate = sample()
                change(candidate)
                with self.assertRaises(ContentError):
                    validate_content(candidate)

    def test_schema_version_must_be_supported_integer(self):
        for version in (True, "1", 2):
            with self.subTest(version=version):
                with self.assertRaises(ContentError):
                    validate_content({"schema_version": version, "name": "Sample"})

    def test_empty_optional_sections_are_accepted_and_not_filled(self):
        candidate = {"schema_version": 1, "name": "Sample", "achievements": [], "education": []}
        self.assertEqual(validate_content(candidate), candidate)
        self.assertNotIn("summary", candidate)

    def test_style_violations_require_approved_input(self):
        for value in ("A passionate assistant.", "Responsible for stock.", "Helped with stock.",
                      "A results-driven assistant.", "A hardworking assistant.", "A dynamic assistant.",
                      "Organized stock records.", "Stock — records."):
            with self.subTest(value=value):
                with self.assertRaises(ContentError):
                    validate_content({"schema_version": 1, "name": "Sample", "summary": value})

    def test_official_titles_and_organisation_names_are_not_rewritten(self):
        candidate = sample()
        candidate["experience"][0]["title"] = "Dynamic Systems Officer"
        candidate["experience"][0]["company"] = "Fictional Center Ltd."
        self.assertEqual(validate_content(candidate)["experience"][0]["title"], "Dynamic Systems Officer")

    def test_control_characters_and_invalid_entry_types_are_rejected(self):
        for value in ("Sample\nCandidate", "Sample\tCandidate", "Sample\x00Candidate", 12, " "):
            with self.subTest(value=value):
                with self.assertRaises(ContentError):
                    validate_content({"schema_version": 1, "name": value})
        candidate = sample()
        candidate["experience"][0]["bullets"] = "not an array"
        with self.assertRaises(ContentError):
            validate_content(candidate)

    def test_duplicate_or_invalid_json_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "candidate.json"
            for content in ('{"name":"A","name":"B","schema_version":1}', '{broken'):
                path.write_text(content, encoding="utf-8")
                with self.assertRaises(ContentError):
                    load_content(path)

    def test_every_supplied_text_fragment_is_available_for_export_checks(self):
        candidate = sample()
        fragments = list(content_fragments(candidate))
        self.assertIn(candidate["experience"][0]["title"], fragments)
        self.assertIn(candidate["experience"][0]["dates"], fragments)
        self.assertIn(candidate["skills"][0]["items"][0], fragments)
        self.assertNotIn(1, fragments)


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.path = Path(self.temporary.name) / "sample.docx"
        self.profile = load_profile()
        self.stage = next(fitting_stages(self.profile))
        self.layout = build_document(sample(), self.path, "Calibri", self.stage)
        self.document = Document(self.path)

    def test_explicit_a4_and_default_margins(self):
        section = self.document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=0.02)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=0.02)
        for side, expected in (("top", .55), ("bottom", .55), ("left", .7), ("right", .7)):
            self.assertAlmostEqual(getattr(section, side + "_margin").inches, expected, delta=.001)

    def test_body_sizes_and_explicit_9_5_point_exceptions(self):
        styles = self.document.styles
        for name in ("Resume Body", "Resume Bullet", "Resume Company", "Resume Entry"):
            self.assertEqual(styles[name].font.size.pt, 10)
        for name in ("Resume Headline", "Resume Contact"):
            self.assertEqual(styles[name].font.size.pt, 9.5)
        self.assertEqual(styles["Resume Name"].font.size.pt, 16)
        self.assertEqual(styles["Resume Heading"].font.size.pt, 10.5)
        for paragraph in self.document.paragraphs:
            if paragraph.style.name == "Resume Entry":
                date = paragraph.runs[-1]
                self.assertEqual(date.font.size.pt, 9.5)
                self.assertTrue(date.italic)
                self.assertFalse(date.bold)
                self.assertEqual(str(date.font.color.rgb), "595959")

    def test_single_spacing_and_justified_body_styles(self):
        self.assertEqual(self.document.styles["Normal"].paragraph_format.line_spacing, 1)
        for name in ("Resume Body", "Resume Bullet", "Resume Company"):
            self.assertEqual(self.document.styles[name].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        for name in ("Resume Name", "Resume Headline", "Resume Contact", "Resume Heading", "Resume Entry"):
            self.assertEqual(self.document.styles[name].paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_date_tab_matches_current_text_width_and_moves_with_margins(self):
        for stage in fitting_stages(self.profile):
            build_document(sample(), self.path, "Calibri", stage)
            document = Document(self.path)
            section = document.sections[0]
            tab = list(document.styles["Resume Entry"].paragraph_format.tab_stops)[0]
            self.assertEqual(tab.alignment, WD_TAB_ALIGNMENT.RIGHT)
            self.assertEqual(tab.position, section.page_width - section.left_margin - section.right_margin)

    def test_heading_rule_and_orphan_prevention(self):
        style = self.document.styles["Resume Heading"]
        self.assertEqual(style.paragraph_format.space_before.pt, 7)
        self.assertEqual(style.paragraph_format.space_after.pt, 3)
        self.assertTrue(style.paragraph_format.keep_with_next)
        rule = style._element.find(".//" + qn("w:bottom"))
        self.assertEqual(rule.get(qn("w:sz")), "6")
        self.assertEqual(rule.get(qn("w:color")), "000000")

    def test_section_order_is_fixed_and_empty_sections_are_omitted(self):
        self.assertEqual(self.layout["sections"], ["SUMMARY", "EDUCATION", "EXPERIENCE", "SKILLS", "ACHIEVEMENTS", "LANGUAGES"])
        candidate = sample()
        candidate["achievements"] = []
        candidate["education"] = []
        layout = build_document(candidate, self.path, "Calibri", self.stage)
        self.assertNotIn("ACHIEVEMENTS", layout["sections"])
        self.assertNotIn("EDUCATION", layout["sections"])

    def test_real_bullets_and_required_indentation(self):
        style = self.document.styles["Resume Bullet"]
        self.assertAlmostEqual(style.paragraph_format.left_indent.inches, .22, delta=1 / 1440)
        self.assertAlmostEqual(style.paragraph_format.first_line_indent.inches, -.18, delta=1 / 1440)
        self.assertEqual(style.paragraph_format.space_after.pt, 1.8)
        for paragraph in self.document.paragraphs:
            if paragraph.style.name == "Resume Bullet":
                self.assertIsNotNone(paragraph._p.find(".//" + qn("w:numId")))
                self.assertFalse(paragraph.text.startswith("•"))
        self.assertTrue(any(node.get(qn("w:val")) == "•" for node in self.document.part.numbering_part.element.iter(qn("w:lvlText"))))

    def test_header_separators_and_skill_label_bold(self):
        candidate = sample()
        paragraphs = self.document.paragraphs
        self.assertEqual(paragraphs[1].text, "  |  ".join(candidate["headline"]))
        self.assertEqual(paragraphs[2].text, "  ·  ".join(candidate["contact"]))
        skill = next(p for p in paragraphs if p.text.startswith("Inventory:"))
        self.assertTrue(skill.runs[0].bold)
        self.assertEqual(skill.text, "Inventory: Stock reconciliation · Physical counts · Delivery checks")

    def test_core_content_is_not_in_tables_headers_or_shapes(self):
        self.assertEqual(len(self.document.tables), 0)
        self.assertEqual(len(self.document.inline_shapes), 0)
        for section in self.document.sections:
            self.assertFalse(any(p.text for p in section.header.paragraphs))
            self.assertFalse(any(p.text for p in section.footer.paragraphs))
        self.assertEqual(self.document.core_properties.author, "")

    def test_fitting_reduces_spacing_before_margins_and_never_below_floor(self):
        stages = list(fitting_stages(self.profile))
        self.assertEqual([s.spacing for s in stages[:4]], [1, .75, .5, 0])
        self.assertTrue(all(s.margins == self.profile["margins_inches"] for s in stages[:4]))
        self.assertTrue(all(min(s.margins.values()) >= .35 for s in stages))
        self.assertTrue(all(s.spacing == 0 for s in stages[4:]))


class ExportTests(unittest.TestCase):
    def test_missing_calibri_cannot_silently_substitute(self):
        with patch.object(builder, "command", return_value="DejaVu Sans"):
            with self.assertRaises(builder.BuildError):
                builder.check_font("Calibri", {"fc-match": "fc-match"})
        with patch.object(builder, "command", return_value="Carlito"):
            builder.check_font("Carlito", {"fc-match": "fc-match"})

    def test_pdf_font_report_requires_requested_family_and_embedding(self):
        header = "name type encoding emb sub uni object ID\n-----------------------\n"
        for row in ("AAAAAA+OtherFont TrueType WinAnsi yes yes yes 9 0", "AAAAAA+Carlito TrueType WinAnsi no yes yes 9 0"):
            with patch.object(builder, "command", return_value=header + row):
                with self.assertRaises(builder.BuildError):
                    builder.inspect_fonts(Path("sample.pdf"), "Carlito", {"pdffonts": "pdffonts"})
        with patch.object(builder, "command", return_value=header + "AAAAAA+Carlito-Bold TrueType WinAnsi yes yes yes 9 0"):
            fonts, _ = builder.inspect_fonts(Path("sample.pdf"), "Carlito", {"pdffonts": "pdffonts"})
            self.assertTrue(fonts[0]["embedded"])

    def test_pdf_text_checks_count_repeated_content(self):
        candidate = {"schema_version": 1, "name": "Sample", "achievements": ["Checked stock", "Checked stock"]}
        with patch.object(builder, "command", return_value="Sample\nACHIEVEMENTS\nChecked stock"):
            with self.assertRaises(builder.BuildError):
                builder.inspect_text(Path("sample.pdf"), candidate, ["ACHIEVEMENTS"], {"pdftotext": "pdftotext"})

    def test_pdf_section_order_must_match_profile(self):
        with patch.object(builder, "command", return_value="Sample\nEXPERIENCE\nEDUCATION"):
            with self.assertRaises(builder.BuildError):
                builder.inspect_text(Path("sample.pdf"), {"name": "Sample"}, ["EDUCATION", "EXPERIENCE"], {"pdftotext": "pdftotext"})

    def test_displaced_dates_are_reported_for_review(self):
        candidate = {"experience": [{"title": "Assistant", "dates": "2024", "company": "Fictional Co."}]}
        self.assertEqual(builder.entry_grouping("Assistant 2024 Fictional Co.", candidate), "passed")
        self.assertEqual(builder.entry_grouping("Assistant Fictional Co. Other content 2024", candidate), "requires_review")
        self.assertEqual(builder.entry_grouping("Sample", {"name": "Sample"}), "not_applicable")

    def test_geometry_rejects_date_wrap_and_text_outside_margins(self):
        candidate = {"experience": [{"title": "Assistant", "dates": "2024"}]}
        layout = {"margins_inches": load_profile()["margins_inches"]}
        for y, x, expected in ((100, 50.5, []), (120, 50.5, ["entry_title_date_wrap_or_collision"]),
                               (100, 0, ["text_outside_margins"])):
            xml = f'<html xmlns="http://www.w3.org/1999/xhtml"><page width="595.276" height="841.89"><word xMin="{x}" xMax="100" yMin="100" yMax="110">Assistant</word><word xMin="520" xMax="544" yMin="{y}" yMax="{y + 10}">2024</word></page></html>'
            with self.subTest(y=y, x=x), patch.object(builder, "command", return_value=xml):
                self.assertEqual(builder.inspect_geometry(Path("sample.pdf"), candidate, layout, {"pdftotext": "pdftotext"}), expected)

    def test_fit_policy_uses_rendered_page_count(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(builder, "build_document", return_value={}), \
             patch.object(builder, "export_pdf", return_value=Path("sample.pdf")), \
             patch.object(builder, "inspect_geometry", return_value=[]), \
             patch.object(builder, "page_count", side_effect=[2, 2, 1]):
            _, _, pages, _, attempts = builder.fit_resume(sample(), Path(directory), "Carlito", {})
            self.assertEqual(pages, 1)
            self.assertEqual(len(attempts), 3)

    def test_overflow_requires_explicit_multiple_page_permission(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(builder, "build_document", return_value={}), \
             patch.object(builder, "export_pdf", return_value=Path("sample.pdf")), \
             patch.object(builder, "inspect_geometry", return_value=[]), \
             patch.object(builder, "page_count", return_value=2):
            with self.assertRaises(builder.OverflowError):
                builder.fit_resume(sample(), Path(directory), "Carlito", {})
            result = builder.fit_resume(sample(), Path(directory), "Carlito", {}, allow_multiple_pages=True)
            self.assertEqual(result[2], 2)

    def test_multiple_page_permission_does_not_allow_horizontal_collisions(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(builder, "build_document", return_value={}), \
             patch.object(builder, "export_pdf", return_value=Path("sample.pdf")), \
             patch.object(builder, "inspect_geometry", return_value=["entry_title_date_wrap_or_collision"]), \
             patch.object(builder, "page_count", return_value=1):
            with self.assertRaises(builder.OverflowError):
                builder.fit_resume(sample(), Path(directory), "Carlito", {}, allow_multiple_pages=True)

    def test_cli_reports_decision_without_claiming_outputs(self):
        output = io.StringIO()
        with patch.object(builder, "build", side_effect=builder.OverflowError([])), contextlib.redirect_stdout(output):
            code = builder.main(["candidate.json"])
        report = json.loads(output.getvalue())
        self.assertEqual(code, 1)
        self.assertEqual(report["status"], "needs_decision")
        self.assertEqual(report["outputs"], [])

    def test_existing_files_and_qa_directories_are_protected(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "existing.pdf"
            output.write_bytes(b"Existing output")
            with self.assertRaises(builder.BuildError):
                builder.build(Path("candidate.json"), output)
            self.assertEqual(output.read_bytes(), b"Existing output")
            with self.assertRaises(builder.BuildError):
                builder.build(Path("candidate.json"), Path(directory) / "new.pdf", qa_dir=Path(directory))


if __name__ == "__main__":
    unittest.main()
