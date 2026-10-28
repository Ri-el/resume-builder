"""Construct the strict A4 profile as a flowing, single-column Word document."""

import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

PROFILE_PATH = Path(__file__).resolve().parents[1] / "profiles/strict-a4.json"


@dataclass(frozen=True)
class FitStage:
    spacing: float
    margins: dict


def load_profile():
    return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))


def fitting_stages(profile):
    for factor in profile["spacing_factors"]:
        yield FitStage(factor, dict(profile["margins_inches"]))
    for margins in profile["compact_margins_inches"]:
        if min(margins.values()) < profile["margin_floor_inches"]:
            raise ValueError("Profile margins are below the permitted floor.")
        yield FitStage(profile["spacing_factors"][-1], dict(margins))


def set_font(target, family, points, colour="000000", bold=False, italic=False):
    target.font.name = family
    target.font.size = Pt(points)
    target.font.bold = bold
    target.font.italic = italic
    target.font.color.rgb = RGBColor.from_string(colour)
    fonts = target._element.get_or_add_rPr().find(qn("w:rFonts"))
    for attribute in list(fonts.attrib):
        del fonts.attrib[attribute]
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + kind), family)
    properties = target._element.get_or_add_rPr()
    for tag, value in (("szCs", str(round(points * 2))), ("bCs", "1" if bold else "0"), ("iCs", "1" if italic else "0")):
        factory = getattr(properties, "get_or_add_" + tag, None)
        node = factory() if factory else properties.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            properties.insert_element_before(node, "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang")
        node.set(qn("w:val"), value)


def create_bullet_numbering(document, family, profile):
    root = document.part.numbering_part.element
    abstract_id = max((int(node.get(qn("w:abstractNumId"))) for node in root.findall(qn("w:abstractNum"))), default=-1) + 1
    num_id = max((int(node.get(qn("w:numId"))) for node in root.findall(qn("w:num"))), default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "•"), ("lvlJc", "left")):
        element = OxmlElement("w:" + tag)
        element.set(qn("w:val"), value)
        level.append(element)
    properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(Inches(profile["bullet"]["left_inches"]).twips))
    tabs.append(tab)
    properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(Inches(profile["bullet"]["left_inches"]).twips))
    indent.set(qn("w:hanging"), str(Inches(profile["bullet"]["hanging_inches"]).twips))
    properties.append(indent)
    level.append(properties)
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), family)
    fonts.set(qn("w:hAnsi"), family)
    run_properties.append(fonts)
    level.append(run_properties)
    abstract.append(level)
    root.append(abstract)
    numbering = OxmlElement("w:num")
    numbering.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    numbering.append(reference)
    root.append(numbering)
    return num_id


def build_document(candidate, destination, family, stage, profile=None):
    profile = profile or load_profile()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = [Mm(value) for value in profile["page_mm"]]
    for side, inches in stage.margins.items():
        setattr(section, side + "_margin", Inches(inches))
    text_width = section.page_width - section.left_margin - section.right_margin
    points = profile["font_points"]
    primary = profile["colours"]["primary"]
    secondary = profile["colours"]["secondary"]
    normal = document.styles["Normal"]
    set_font(normal, family, points["body"], primary)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "en-GB")
    normal._element.get_or_add_rPr().append(language)
    line_height_factor = profile.get("line_height_factor")
    normal.paragraph_format.line_spacing = Pt(points["body"] * line_height_factor) if line_height_factor is not None else 1
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True
    # Suppress automatic hyphenation so extraction is comparable with input.
    suppress = OxmlElement("w:suppressAutoHyphens")
    normal._element.get_or_add_pPr().append(suppress)

    def style(name, size, *, alignment=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, colour=primary):
        value = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        value.base_style = normal
        set_font(value, family, size, colour, bold, italic)
        value.paragraph_format.alignment = alignment
        value.paragraph_format.keep_together = True
        if line_height_factor is not None:
            value.paragraph_format.line_spacing = Pt(size * line_height_factor)
        return value

    style("Resume Name", points["name"], bold=True)
    style("Resume Headline", points["headline"], bold=True)
    style("Resume Contact", points["contact"])
    heading = style("Resume Heading", points["heading"], bold=True)
    heading.paragraph_format.keep_with_next = True
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    heading._element.get_or_add_pPr().append(outline)
    heading.paragraph_format.space_before = Pt(profile["heading"]["before_points"] * stage.spacing)
    heading.paragraph_format.space_after = Pt(profile["heading"]["after_points"] * stage.spacing)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for attribute, value in {"val": "single", "sz": str(round(profile["heading"]["rule_points"] * 8)), "space": "1", "color": primary}.items():
        bottom.set(qn("w:" + attribute), value)
    borders.append(bottom)
    heading._element.get_or_add_pPr().append(borders)
    title = style("Resume Entry", points["title"], bold=True)
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    company = style("Resume Company", points["body"], italic=True, colour=secondary, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    body_style = style("Resume Body", points["body"], alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    bullet_style = style("Resume Bullet", points["body"], alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    bullet_style.paragraph_format.left_indent = Inches(profile["bullet"]["left_inches"])
    bullet_style.paragraph_format.first_line_indent = -Inches(profile["bullet"]["hanging_inches"])
    bullet_style.paragraph_format.space_after = Pt(profile["bullet"]["after_points"] * stage.spacing)
    number_id = create_bullet_numbering(document, family, profile)

    def bullet(value):
        paragraph = document.add_paragraph(value, bullet_style)
        numbering = OxmlElement("w:numPr")
        for tag, val in (("ilvl", "0"), ("numId", str(number_id))):
            node = OxmlElement("w:" + tag)
            node.set(qn("w:val"), val)
            numbering.append(node)
        paragraph._p.get_or_add_pPr().append(numbering)

    def entry(item, title_key, organisation_key, bullets_key):
        paragraph = document.add_paragraph(style=title)
        paragraph.add_run(item[title_key])
        date = paragraph.add_run("\t" + item["dates"])
        set_font(date, family, points["date"], secondary, italic=True)
        organisation = item[organisation_key]
        if item.get("location"):
            organisation += ", " + item["location"]
        line = document.add_paragraph(organisation, company)
        line.paragraph_format.keep_with_next = bool(item.get(bullets_key))
        for value in item.get(bullets_key, []):
            bullet(value)

    document.add_paragraph(candidate["name"], "Resume Name")
    if candidate.get("headline"):
        document.add_paragraph("  |  ".join(candidate["headline"]), "Resume Headline")
    if candidate.get("contact"):
        document.add_paragraph("  ·  ".join(candidate["contact"]), "Resume Contact")
    included_sections = []
    for key in profile["section_order"]:
        values = candidate.get(key)
        if not values:
            continue
        included_sections.append(profile["section_labels"][key])
        document.add_paragraph(profile["section_labels"][key], heading)
        if key == "summary":
            document.add_paragraph(values, body_style)
        elif key == "education":
            for item in values:
                entry(item, "qualification", "institution", "details")
        elif key == "experience":
            for item in values:
                entry(item, "title", "company", "bullets")
        elif key == "skills":
            for item in values:
                paragraph = document.add_paragraph(style=body_style)
                paragraph.add_run(item["category"] + ":").bold = True
                paragraph.add_run(" " + " · ".join(item["items"]))
        elif key == "achievements":
            for value in values:
                bullet(value)
        elif key == "languages":
            for item in values:
                paragraph = document.add_paragraph(style=body_style)
                paragraph.add_run(item["language"] + ":").bold = True
                paragraph.add_run(" " + item["level"])

    for field in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
        setattr(document.core_properties, field, "")
    document.save(destination)
    return {
        "spacing_factor": stage.spacing,
        "line_height_factor": line_height_factor,
        "margins_inches": stage.margins,
        "date_tab_inches": round(text_width / Inches(1), 6),
        "sections": included_sections,
    }
